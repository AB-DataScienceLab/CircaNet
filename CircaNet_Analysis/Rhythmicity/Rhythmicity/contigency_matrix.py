import os
import numpy as np
import pandas as pd
from scipy.stats import fisher_exact
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

# ---------------------------------------------------------------
# Settings
# ---------------------------------------------------------------
CYCLOPS_FILE = "cyclops_hgnc_mapped.csv"
ALL_FILE = "All_Protein.tsv"
CIRC_FILE = "Circadian_Protein.tsv"
Q_CUTOFF = 0.05
OUT_DIR = "tissue_workbooks"
os.makedirs(OUT_DIR, exist_ok=True)

# ---------------------------------------------------------------
# Load + classify (same logic as adipose_hgnc.py)
# ---------------------------------------------------------------
data = pd.read_csv(CYCLOPS_FILE)
all_genes = pd.read_csv(ALL_FILE, sep="\t")
circadian = pd.read_csv(CIRC_FILE, sep="\t")

for df in [data, all_genes, circadian]:
    df["hgnc_id"] = df["hgnc_id"].astype("string").str.strip()

all_ids = set(all_genes["hgnc_id"].dropna())
circadian_ids = set(circadian["hgnc_id"].dropna())

data["FDR_q"] = pd.to_numeric(data["FDR_q"], errors="coerce")
matched = data[data["hgnc_id"].notna() & data["hgnc_id"].isin(all_ids)].copy()
matched = matched.dropna(subset=["FDR_q"])

matched["classification"] = np.where(
    matched["hgnc_id"].isin(circadian_ids), "Circadian", "Non-circadian")
matched["CYCLOPS_result"] = np.where(
    matched["FDR_q"] <= Q_CUTOFF, "q <= 0.05", "q > 0.05")

ROWS = ["Circadian", "Non-circadian"]
COLS = ["q <= 0.05", "q > 0.05"]

# ---------------------------------------------------------------
# Counts + Fisher per tissue
# ---------------------------------------------------------------
records = []
for tissue, sub in matched.groupby("tissue"):
    m = pd.crosstab(sub["classification"], sub["CYCLOPS_result"]).reindex(
        index=ROWS, columns=COLS, fill_value=0)
    a, b = int(m.loc["Circadian", COLS[0]]), int(m.loc["Circadian", COLS[1]])
    c, d = int(m.loc["Non-circadian", COLS[0]]), int(m.loc["Non-circadian", COLS[1]])
    try:
        _, p = fisher_exact([[a, b], [c, d]], alternative="two-sided")
    except ValueError:
        p = np.nan
    records.append({"tissue": tissue, "a": a, "b": b, "c": c, "d": d, "p": p})

res = pd.DataFrame(records)


def bh(p):
    p = np.asarray(p, dtype=float)
    out = np.full_like(p, np.nan)
    ok = ~np.isnan(p)
    pv = p[ok]
    order = np.argsort(pv)
    ranked = pv[order] * len(pv) / (np.arange(len(pv)) + 1)
    ranked = np.minimum.accumulate(ranked[::-1])[::-1]
    adj = np.empty_like(pv)
    adj[order] = np.minimum(ranked, 1)
    out[ok] = adj
    return out


res["p_bh"] = bh(res["p"])

# ---------------------------------------------------------------
# Workbook styling
# ---------------------------------------------------------------
F = "Arial"
bold = Font(name=F, bold=True, size=11)
norm = Font(name=F, size=11)
title = Font(name=F, bold=True, size=14)
hdr_fill = PatternFill("solid", start_color="D9E1F2")
tot_fill = PatternFill("solid", start_color="F2F2F2")
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")


def put(ws, ref, value, font=norm, fill=None, border=True, align=center, fmt=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    cell.alignment = align
    if fill:
        cell.fill = fill
    if border:
        cell.border = box
    if fmt:
        cell.number_format = fmt
    return cell


for r in res.itertuples():
    wb = Workbook()
    ws = wb.active
    ws.title = "Confusion_Matrix"

    ws["A1"].value = f"{r.tissue}: circadian classification vs CYCLOPS significance"
    ws["A1"].font = title

    # Matrix (rows 3-6, columns A-D)
    put(ws, "A3", "Classification", bold, hdr_fill)
    put(ws, "B3", "q <= 0.05", bold, hdr_fill)
    put(ws, "C3", "q > 0.05", bold, hdr_fill)
    put(ws, "D3", "Row total", bold, hdr_fill)

    put(ws, "A4", "Circadian", bold)
    put(ws, "B4", r.a, fmt="#,##0")
    put(ws, "C4", r.b, fmt="#,##0")
    put(ws, "D4", "=B4+C4", fill=tot_fill, fmt="#,##0")

    put(ws, "A5", "Non-circadian", bold)
    put(ws, "B5", r.c, fmt="#,##0")
    put(ws, "C5", r.d, fmt="#,##0")
    put(ws, "D5", "=B5+C5", fill=tot_fill, fmt="#,##0")

    put(ws, "A6", "Column total", bold, tot_fill)
    put(ws, "B6", "=B4+B5", fill=tot_fill, fmt="#,##0")
    put(ws, "C6", "=C4+C5", fill=tot_fill, fmt="#,##0")
    put(ws, "D6", "=D4+D5", bold, tot_fill, fmt="#,##0")

    # Statistics block (a=B4, b=C4, c=B5, d=C5)
    L = Alignment(horizontal="left", vertical="center", wrap_text=True)
    put(ws, "A8", "Parameter", bold, hdr_fill)
    put(ws, "B8", "Value", bold, hdr_fill)
    put(ws, "C8", "Definition", bold, hdr_fill)
    ws.merge_cells("C8:D8")

    put(ws, "A9", "Circadian significant fraction", bold, align=L)
    put(ws, "B9", "=IF(D4=0,\"NA\",B4/D4)", fmt="0.00%")
    put(ws, "C9", "a/(a+b)", align=L)

    put(ws, "A10", "Non-circadian significant fraction", bold, align=L)
    put(ws, "B10", "=IF(D5=0,\"NA\",B5/D5)", fmt="0.00%")
    put(ws, "C10", "c/(c+d)", align=L)

    put(ws, "A11", "Non-circadian non-significant fraction", bold, align=L)
    put(ws, "B11", "=IF(D5=0,\"NA\",C5/D5)", fmt="0.00%")
    put(ws, "C11", "d/(c+d)", align=L)

    put(ws, "A12", "Odds ratio", bold, align=L)
    put(ws, "B12", "=IF(C4*B5=0,\"NA\",(B4*C5)/(C4*B5))", fmt="0.000")
    put(ws, "C12", "ad/(bc); NA when bc = 0", align=L)

    put(ws, "A13", "Fisher's exact p-value (two-sided)", bold, align=L)
    c13 = put(ws, "B13", None if np.isnan(r.p) else float(r.p), fmt="0.00E+00")
    put(ws, "C13", "Tests whether classification and CYCLOPS significance are associated", align=L)
    c13.comment = Comment(
        "Computed in Python with scipy.stats.fisher_exact "
        "(two-sided). Excel has no exact two-sided Fisher function, "
        "so this is a stored value, not a formula.", "script")

    put(ws, "A14", "BH-adjusted p-value", bold, align=L)
    c14 = put(ws, "B14", None if np.isnan(r.p_bh) else float(r.p_bh), fmt="0.00E+00")
    put(ws, "C14", "Benjamini-Hochberg across all tissues in the run", align=L)
    c14.comment = Comment(
        "Benjamini-Hochberg across the Fisher p-values of all tissues "
        "in the run. Stored value.", "script")

    put(ws, "A15", "Significant after BH (< 0.05)", bold, align=L)
    put(ws, "B15", "=IF(ISNUMBER(B14),IF(B14<0.05,\"Yes\",\"No\"),\"NA\")")
    put(ws, "C15", "BH-adjusted p < 0.05", align=L)

    for rr in (9, 10, 11, 12, 13, 14, 15):
        ws.merge_cells(f"C{rr}:D{rr}")
        ws[f"D{rr}"].border = box
    ws.row_dimensions[13].height = 45
    ws.row_dimensions[14].height = 30

    ws["A17"].value = ("Notes: q-value cutoff = 0.05. Classification = presence of HGNC ID in "
                       "Circadian_Protein.tsv (restricted to genes in All_Protein.tsv). "
                       "Counts are stored values from the CYCLOPS output; totals, fractions "
                       "and odds ratio are formulas.")
    ws["A17"].font = Font(name=F, italic=True, size=9)
    ws["A17"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A17:D19")

    ws.column_dimensions["A"].width = 40
    for col in "BCD":
        ws.column_dimensions[col].width = 16

    wb.save(os.path.join(OUT_DIR, f"{r.tissue}_confusion_fisher.xlsx"))

res["circadian_sig_fraction"] = res["a"] / (res["a"] + res["b"])
res["noncirc_sig_fraction"] = res["c"] / (res["c"] + res["d"])
res["noncirc_nonsig_fraction"] = res["d"] / (res["c"] + res["d"])
res["odds_ratio"] = np.where(res["b"] * res["c"] > 0,
                             (res["a"] * res["d"]) / (res["b"] * res["c"]).replace(0, np.nan),
                             np.nan)
res.rename(columns={"a": "circ_sig", "b": "circ_nonsig", "c": "noncirc_sig",
                    "d": "noncirc_nonsig", "p": "fisher_p", "p_bh": "fisher_p_BH"}
           ).to_csv(os.path.join(OUT_DIR, "_summary_all_tissues.csv"), index=False)

print(f"Created {len(res)} workbooks in ./{OUT_DIR}/")
